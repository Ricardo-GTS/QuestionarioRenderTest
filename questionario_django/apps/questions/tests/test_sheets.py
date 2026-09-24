"""Planilha de backup das questoes (formato do Google Forms original), uma por semestre,
regenerada a partir do banco a cada mudanca; e planilha por periodo para o admin."""

from datetime import date, datetime
from datetime import timezone as dt_timezone
from io import BytesIO, StringIO
from pathlib import Path
from unittest import mock

import pytest
from django.core.management import call_command
from django.db import connection
from django.urls import reverse
from django_tenants.utils import schema_context
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.core.models import Semester
from apps.questions import sheets
from apps.questions.models import Question, QuestionStatus
from conftest import _fake_get_embedding

ORIGINAL_COLUMNS = [
    "Timestamp", "Email Address", "Nome Completo", "Tópico da questão", "Questão",
    "Resposta", "Citações e referências", "Pertinência",
]


def _rows(path):
    sheet = load_workbook(path).active
    assert sheet.title == "Form Responses 1"
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def _active():
    return Semester.objects.get(is_active=True)


def _question(author, statement, topic="Geral", correct=True, status=QuestionStatus.ACTIVE):
    return Question.objects.create(
        author=author, statement=statement, correct_answer=correct, topic=topic,
        citations_references="Livro X", pertinence="Relevante", embedding=_fake_get_embedding(statement), status=status,
    )


@pytest.fixture
def author():
    return User.objects.create_user(email="autor@example.com", name="Autor Um", password="x" * 8, registration_number="10000001")


@pytest.fixture
def admin_client(client):
    admin = User.objects.create_user(email="admin@example.com", name="Admin", password="x" * 8, registration_number="99990000")
    client.force_login(admin)
    return client


# --- puras ---


def test_question_row_and_columns():
    q = Question(
        statement="S", correct_answer=False, topic="T", citations_references="C", pertinence="P",
        status="reported", created_at=datetime(2026, 9, 24, 2, 30, tzinfo=dt_timezone.utc),
    )
    q.author = User(email="a@example.com", name="Fulano")
    row = sheets.question_row(q)
    assert row[0] == datetime(2026, 9, 23, 23, 30)  # Recife, sem tzinfo
    assert row[1:] == ["a@example.com", "Fulano", "T", "S", "Falsa", "C", "P", "Em análise"]
    assert sheets.COLUMNS[:8] == ORIGINAL_COLUMNS and sheets.COLUMNS[8] == "Situação"


def test_parse_range():
    available = (date(2026, 5, 1), date(2026, 9, 24))
    assert sheets.parse_range("2026-06-10", "2026-06-01", available)[1].startswith("A data de início")
    assert "Escolha datas entre" in sheets.parse_range("2026-04-01", "2026-06-01", available)[1]
    assert sheets.parse_range("abc", "2026-06-01", available)[1] == "Informe datas válidas."
    assert sheets.parse_range("2026-06-01", "2026-06-30", None)[1]
    (start, end), error = sheets.parse_range("2026-06-01", "2026-06-30", available)
    assert error is None
    assert start.isoformat() == "2026-06-01T00:00:00-03:00" and end.isoformat() == "2026-07-01T00:00:00-03:00"


# --- planilha do semestre ---


@pytest.mark.django_db
def test_creating_question_writes_sheet(client, author, django_capture_on_commit_callbacks, fake_embedding):
    client.force_login(author)
    with django_capture_on_commit_callbacks(execute=True):
        client.post(
            reverse("questions:create"),
            {"statement": "O Sol e uma estrela", "correct_answer": "true", "topic": "__new_topic__",
             "new_topic": "Astronomia", "citations_references": "NASA", "pertinence": "Basico"},
        )
    rows = _rows(sheets.sheet_path(_active()))
    assert rows[0] == ORIGINAL_COLUMNS + ["Situação"]
    assert rows[1][1:] == ["autor@example.com", "Autor Um", "Astronomia", "O Sol e uma estrela", "Verdadeira", "NASA", "Basico", "Ativa"]
    assert isinstance(rows[1][0], datetime)


@pytest.mark.django_db
def test_sheet_follows_edits_status_topics_and_deletes(author, django_capture_on_commit_callbacks):
    from apps.moderation.services import delete_topic, rename_topic, resolve_reported_question, update_question

    with django_capture_on_commit_callbacks(execute=True):
        q1 = _question(author, "Primeira", topic="A")
        q2 = _question(author, "Segunda", topic="B")
    path = sheets.sheet_path(_active())
    assert [r[4] for r in _rows(path)[1:]] == ["Primeira", "Segunda"]

    with django_capture_on_commit_callbacks(execute=True):
        update_question(q1, statement="Primeira editada")
    assert _rows(path)[1][4] == "Primeira editada"

    with django_capture_on_commit_callbacks(execute=True):
        resolve_reported_question(q2, approve_removal=True)
    assert _rows(path)[2][8] == "Removida"

    with django_capture_on_commit_callbacks(execute=True):
        rename_topic("A", "A renomeado")  # .update(): chamada explicita
    assert _rows(path)[1][3] == "A renomeado"

    with django_capture_on_commit_callbacks(execute=True):
        delete_topic("B", delete_questions=True)
    assert [r[4] for r in _rows(path)[1:]] == ["Primeira editada"]


@pytest.mark.django_db
def test_one_file_per_semester(author, django_capture_on_commit_callbacks):
    from apps.core.semesters import open_new_semester

    with django_capture_on_commit_callbacks(execute=True):
        _question(author, "Do primeiro semestre")
    old = _active()
    old_rows = _rows(sheets.sheet_path(old))
    connection.set_schema_to_public()
    new = open_new_semester(name="2026.2", copy_topics=False, backup_done=True)
    with schema_context(new.schema_name), django_capture_on_commit_callbacks(execute=True):
        _question(author, "Do segundo semestre")
    assert [r[4] for r in _rows(sheets.sheet_path(new))[1:]] == ["Do segundo semestre"]
    assert _rows(sheets.sheet_path(old)) == old_rows
    assert sheets.sheet_path(new).name == "Banco_de_Questoes_2026.2.xlsx"


@pytest.mark.django_db
def test_sheet_failure_does_not_block_question(client, author, django_capture_on_commit_callbacks, fake_embedding, caplog):
    client.force_login(author)
    with mock.patch("apps.questions.sheets.build_workbook", side_effect=OSError("disco cheio")):
        with django_capture_on_commit_callbacks(execute=True):
            resp = client.post(
                reverse("questions:create"),
                {"statement": "Salva mesmo assim", "correct_answer": "true", "topic": "__new_topic__",
                 "new_topic": "X", "citations_references": "r", "pertinence": "p"},
            )
    assert resp.context["success_message"]
    assert Question.objects.filter(statement="Salva mesmo assim").exists()
    assert "Falha ao gravar a planilha" in caplog.text


@pytest.mark.django_db
def test_atomic_replace_keeps_previous_file_on_failure(author, django_capture_on_commit_callbacks):
    with django_capture_on_commit_callbacks(execute=True):
        _question(author, "Original")
    path = sheets.sheet_path(_active())
    before = path.read_bytes()
    with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("queda no meio")):
        with pytest.raises(OSError):
            sheets.write_sheet(_active())
    assert path.read_bytes() == before  # arquivo anterior intacto
    assert not list(Path(path.parent).glob(".*.tmp"))  # temporario limpo


@pytest.mark.django_db
def test_repeated_regeneration_is_complete(author):
    for i in range(3):
        _question(author, f"Q{i}")
    semester = _active()
    sheets.write_sheet(semester)
    sheets.write_sheet(semester)
    assert [r[4] for r in _rows(sheets.sheet_path(semester))[1:]] == ["Q0", "Q1", "Q2"]


@pytest.mark.django_db
def test_async_worker_coalesces(author, settings, monkeypatch):
    """Duas mudancas enquanto uma gravacao roda: grava de novo UMA vez no fim."""
    settings.QUESTION_SHEETS_ASYNC = True
    calls = []
    semester = _active()

    def fake_write(sid):
        calls.append(sid)
        if len(calls) == 1:  # durante a 1a gravacao chegam mais 2 mudancas
            sheets._kick(sid)
            sheets._kick(sid)

    started = []
    monkeypatch.setattr(sheets, "_write_safely", fake_write)
    monkeypatch.setattr(sheets.connections, "close_all", lambda: None)
    monkeypatch.setattr(sheets.threading, "Thread", lambda target, args, daemon: type("T", (), {"start": lambda self: (started.append(1), target(*args))})())
    sheets._kick(semester.pk)
    assert calls == [semester.pk, semester.pk]  # 1a + uma rodada extra (nao 3)
    assert len(started) == 1


# --- telas do admin ---


@pytest.mark.django_db
def test_download_and_regenerate_require_admin(client, author):
    client.force_login(author)
    semester = _active()
    assert client.get(reverse("core:download_sheet", args=[semester.id])).status_code == 403
    assert client.post(reverse("core:regenerate_sheet", args=[semester.id])).status_code == 403
    assert client.post(reverse("core:export_range"), {"start": "2026-01-01", "end": "2026-12-31"}).status_code == 403


@pytest.mark.django_db
def test_download_generates_when_missing_and_regenerate(admin_client, author):
    _question(author, "Sem arquivo ainda")  # sem on_commit: o arquivo nao existe
    semester = _active()
    assert not sheets.sheet_path(semester).exists()
    resp = admin_client.get(reverse("core:download_sheet", args=[semester.id]))
    assert resp.status_code == 200
    assert 'filename="Banco_de_Questoes_2026.1.xlsx"' in resp["Content-Disposition"]
    body = b"".join(resp.streaming_content)
    assert load_workbook(BytesIO(body)).active.cell(2, 5).value == "Sem arquivo ainda"
    resp = admin_client.post(reverse("core:regenerate_sheet", args=[semester.id]))
    assert resp.status_code == 302


@pytest.mark.django_db
def test_gerar_planilha_command(author):
    _question(author, "Via comando")
    call_command("gerar_planilha", semestre="2026.1", stdout=StringIO())
    assert [r[4] for r in _rows(sheets.sheet_path(_active()))[1:]] == ["Via comando"]


# --- planilha por periodo ---


def _at(question, when):
    Question.objects.filter(pk=question.pk).update(created_at=when)


@pytest.mark.django_db
def test_range_export_across_semesters(admin_client, author):
    from apps.core.semesters import open_new_semester

    q_old = _question(author, "Junho no 2026.1")
    _at(q_old, datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc))
    late = _question(author, "Ultimo dia 23h30 local")
    _at(late, datetime(2026, 7, 1, 2, 30, tzinfo=dt_timezone.utc))  # 30/06 23h30 em Recife
    outside = _question(author, "Julho, fora")
    _at(outside, datetime(2026, 7, 15, 12, 0, tzinfo=dt_timezone.utc))
    connection.set_schema_to_public()
    new = open_new_semester(name="2026.2", copy_topics=False, backup_done=True)
    with schema_context(new.schema_name):
        q_new = _question(author, "Junho no 2026.2")
        _at(q_new, datetime(2026, 6, 20, 12, 0, tzinfo=dt_timezone.utc))

    assert sheets.available_range() == (date(2026, 6, 10), date(2026, 7, 15))
    page = admin_client.get(reverse("core:semesters"))
    assert "10/06/2026" in page.content.decode() and "15/07/2026" in page.content.decode()

    # o disponivel comeca em 10/06 (primeira questao); datas antes disso sao recusadas
    resp = admin_client.post(reverse("core:export_range"), {"start": "2026-06-10", "end": "2026-06-30"})
    assert resp.status_code == 200
    assert 'filename="Banco_de_Questoes_2026-06-10_a_2026-06-30.xlsx"' in resp["Content-Disposition"]
    rows = [list(r) for r in load_workbook(BytesIO(resp.content)).active.iter_rows(values_only=True)]
    assert rows[0] == ORIGINAL_COLUMNS + ["Situação", "Semestre"]
    assert [(r[4], r[9]) for r in rows[1:]] == [
        ("Junho no 2026.1", "2026.1"),
        ("Junho no 2026.2", "2026.2"),
        ("Ultimo dia 23h30 local", "2026.1"),
    ]


@pytest.mark.django_db
def test_range_export_errors_and_empty(admin_client, author):
    first, last = _question(author, "Inicio"), _question(author, "Fim")
    _at(first, datetime(2026, 6, 1, 12, 0, tzinfo=dt_timezone.utc))
    _at(last, datetime(2026, 6, 30, 12, 0, tzinfo=dt_timezone.utc))
    resp = admin_client.post(reverse("core:export_range"), {"start": "2026-06-20", "end": "2026-06-10"})
    assert "A data de início" in resp.context["error"]
    assert resp["Content-Type"].startswith("text/html")
    # periodo valido, dentro do disponivel, mas sem nenhuma questao: aviso, sem arquivo
    resp = admin_client.post(reverse("core:export_range"), {"start": "2026-06-10", "end": "2026-06-20"})
    assert resp.context["error"] == "Nenhuma questão nesse período."
    assert resp["Content-Type"].startswith("text/html")
    # um dia so', com a questao: baixa
    resp = admin_client.post(reverse("core:export_range"), {"start": "2026-06-01", "end": "2026-06-01"})
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")


# --- regressao: semestre certo dentro de schema_context() ---


@pytest.mark.django_db
def test_connection_semester_inside_schema_context(author, django_capture_on_commit_callbacks):
    """Dentro de schema_context() o django-tenants poe um FakeTenant na conexao; a planilha
    e as configuracoes tem que usar o semestre DAQUELE schema, nao o ativo."""
    from apps.core.semesters import connection_semester, open_new_semester
    from apps.core.services import get_effective_settings

    old = _active()
    connection.set_schema_to_public()
    new = open_new_semester(name="2026.2", copy_topics=False, backup_done=True)
    Semester.objects.filter(pk=old.pk).update(quiz_size=7)
    with schema_context(old.schema_name):
        assert connection_semester().pk == old.pk
        assert get_effective_settings().quiz_size == 7  # do semestre encerrado, nao do ativo
        with django_capture_on_commit_callbacks(execute=True):
            _question(author, "Criada no semestre encerrado")
    assert [r[4] for r in _rows(sheets.sheet_path(old))[1:]] == ["Criada no semestre encerrado"]
    assert not sheets.sheet_path(new).exists()


# --- compatibilidade com o importador (pasta analise, local) ---


@pytest.mark.django_db
def test_generated_sheet_is_readable_by_importer(author):
    from django.apps import apps as django_apps

    if not django_apps.is_installed("analise"):
        pytest.skip("pasta analise (local) nao instalada")
    from analise.management.commands.import_question_bank import Command as ImportCommand

    q = _question(author, "Ida e volta", topic="Compat", correct=False)
    path = sheets.write_sheet(_active())
    [row] = ImportCommand()._read_rows(str(path))
    assert (row["email"], row["name"], row["topic"], row["statement"], row["correct_answer"]) == (
        "autor@example.com", "Autor Um", "Compat", "Ida e volta", False,
    )
    assert abs((row["timestamp"] - q.created_at).total_seconds()) < 1  # mesmo instante
